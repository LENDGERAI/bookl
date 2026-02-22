from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader

from booklify.core.country_rules import CountryProfile


@dataclass(slots=True)
class ExtractedTransaction:
    tx_type: str | None
    category: str | None
    description: str | None
    counterparty: str | None
    reference_number: str | None
    occurred_on: date | None
    amount: Decimal
    tax_amount: Decimal
    currency: str | None
    payment_method: str | None
    line_items: list[dict[str, Any]] = field(default_factory=list)
    confidence_hint: float = 0.5
    raw: dict[str, Any] = field(default_factory=dict)


class ExtractionService:
    def extract(self, path: Path, country_profile: CountryProfile) -> list[ExtractedTransaction]:
        ext = path.suffix.lower()
        if ext == ".csv":
            return self._extract_csv(path, country_profile)
        if ext == ".xlsx":
            return self._extract_xlsx(path, country_profile)
        if ext == ".pdf":
            return self._extract_pdf(path, country_profile)
        if ext in {".png", ".jpg", ".jpeg", ".txt", ".xls"}:
            return self._extract_textlike(path, country_profile)
        return []

    @staticmethod
    def _safe_decimal(value: Any, default: Decimal = Decimal("0")) -> Decimal:
        if value is None:
            return default
        if isinstance(value, Decimal):
            return value
        cleaned = str(value).strip().replace(",", "")
        if not cleaned:
            return default
        cleaned = re.sub(r"[^0-9.\-]", "", cleaned)
        if not cleaned or cleaned == "-":
            return default
        try:
            return Decimal(cleaned)
        except InvalidOperation:
            return default

    @staticmethod
    def _parse_date(value: Any) -> date | None:
        if not value:
            return None
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        raw = str(value).strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%Y/%m/%d"):
            try:
                return datetime.strptime(raw, fmt).date()
            except ValueError:
                continue
        try:
            return datetime.fromisoformat(raw).date()
        except ValueError:
            return None

    @staticmethod
    def _map_tx_type(raw_type: str | None, amount: Decimal, description: str | None) -> str:
        value = (raw_type or "").strip().lower()
        if value in {"income", "expense", "receivable", "payable", "payroll", "marketing", "tax"}:
            return value
        desc = (description or "").lower()
        if "invoice due" in desc or "receivable" in desc:
            return "receivable"
        if "bill due" in desc or "payable" in desc:
            return "payable"
        if "salary" in desc or "payroll" in desc:
            return "payroll"
        if "ads" in desc or "marketing" in desc:
            return "marketing"
        return "income" if amount >= 0 else "expense"

    def _row_to_transaction(self, row: dict[str, Any], country_profile: CountryProfile) -> ExtractedTransaction | None:
        lower = {str(k).strip().lower(): v for k, v in row.items() if k is not None}
        description = self._first(
            lower,
            "description",
            "details",
            "memo",
            "narration",
            default="",
        )
        amount = self._safe_decimal(self._first(lower, "amount", "total", "gross", default=Decimal("0")))
        if amount == 0:
            return None
        tax_amount = self._safe_decimal(self._first(lower, "tax", "vat", "gst", default=Decimal("0")))
        tx_type = self._map_tx_type(self._first(lower, "type", "transaction_type"), amount, description)

        return ExtractedTransaction(
            tx_type=tx_type,
            category=self._first(lower, "category"),
            description=description or None,
            counterparty=self._first(lower, "vendor", "client", "counterparty", "name"),
            reference_number=self._first(lower, "invoice_number", "receipt_number", "reference", "ref", "number"),
            occurred_on=self._parse_date(self._first(lower, "date", "issued_on", "invoice_date")),
            amount=abs(amount),
            tax_amount=tax_amount,
            currency=(self._first(lower, "currency", default=country_profile.currency) or country_profile.currency).upper(),
            payment_method=self._first(lower, "payment_method", "method"),
            line_items=[],
            confidence_hint=0.88,
            raw=lower,
        )

    @staticmethod
    def _first(source: dict[str, Any], *keys: str, default: Any = None) -> Any:
        for key in keys:
            if key in source and source[key] not in (None, ""):
                return source[key]
        return default

    def _extract_csv(self, path: Path, country_profile: CountryProfile) -> list[ExtractedTransaction]:
        content = path.read_text(encoding="utf-8", errors="ignore")
        reader = csv.DictReader(io.StringIO(content))
        results: list[ExtractedTransaction] = []
        for row in reader:
            tx = self._row_to_transaction(row, country_profile)
            if tx:
                results.append(tx)
        return results

    def _extract_xlsx(self, path: Path, country_profile: CountryProfile) -> list[ExtractedTransaction]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value).strip() if value is not None else "" for value in rows[0]]
        results: list[ExtractedTransaction] = []
        for item in rows[1:]:
            record = {headers[idx]: item[idx] for idx in range(min(len(headers), len(item)))}
            tx = self._row_to_transaction(record, country_profile)
            if tx:
                results.append(tx)
        return results

    def _extract_pdf(self, path: Path, country_profile: CountryProfile) -> list[ExtractedTransaction]:
        reader = PdfReader(str(path))
        combined = "\n".join(page.extract_text() or "" for page in reader.pages)
        return self._extract_from_text(combined, country_profile)

    def _extract_textlike(self, path: Path, country_profile: CountryProfile) -> list[ExtractedTransaction]:
        text = path.read_text(encoding="utf-8", errors="ignore")
        return self._extract_from_text(text, country_profile, low_confidence=True)

    def _extract_from_text(
        self,
        text: str,
        country_profile: CountryProfile,
        low_confidence: bool = False,
    ) -> list[ExtractedTransaction]:
        if not text.strip():
            return []
        invoice_ref = self._capture(text, r"(?:invoice|receipt)\s*(?:number|no|#)?\s*[:\-]?\s*([A-Z0-9\-\/]+)")
        counterparty = self._capture(text, r"(?:vendor|supplier|from|client)\s*[:\-]\s*([^\n\r]+)")
        payment_method = self._capture(text, r"(?:payment method|paid via)\s*[:\-]\s*([^\n\r]+)")
        currency = self._capture(text, r"\b(AUD|USD|GBP|INR|EUR)\b") or country_profile.currency

        amount = self._safe_decimal(self._capture(text, r"(?:total|amount due|grand total)\s*[:\-]?\s*([0-9,]+\.\d{1,2})"))
        tax_amount = self._safe_decimal(self._capture(text, r"(?:tax|vat|gst)\s*[:\-]?\s*([0-9,]+\.\d{1,2})"))
        date_value = self._capture(
            text,
            r"(\d{4}-\d{2}-\d{2}|\d{2}/\d{2}/\d{4}|\d{2}-\d{2}-\d{4})",
        )

        if amount <= 0:
            amount = Decimal("0")
            candidates = re.findall(r"([0-9,]+\.\d{1,2})", text)
            if candidates:
                amount = self._safe_decimal(candidates[-1], default=Decimal("0"))
        if amount <= 0:
            return []

        desc = " ".join(text.split())[:220]
        tx_type = "income" if "invoice" in text.lower() else "expense"
        confidence = 0.46 if not low_confidence else 0.34
        return [
            ExtractedTransaction(
                tx_type=tx_type,
                category=None,
                description=desc,
                counterparty=counterparty,
                reference_number=invoice_ref,
                occurred_on=self._parse_date(date_value),
                amount=amount,
                tax_amount=tax_amount,
                currency=currency,
                payment_method=payment_method,
                line_items=[],
                confidence_hint=confidence,
                raw={"text_excerpt": desc},
            )
        ]

    @staticmethod
    def _capture(text: str, pattern: str) -> str | None:
        match = re.search(pattern, text, re.IGNORECASE)
        return match.group(1).strip() if match else None

